import base64
import hashlib
import io
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4
from xml.sax.saxutils import escape

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.auth.context import current_actor
from app.database.document_export_models import DocumentExport
from app.database.models import SignatureRecord
from app.database.session import begin_sqlite_immediate
from app.domain.errors import DomainError
from app.repositories import audit as audit_repo, documents as document_repo
from app.services.case_management_service import assert_case_not_void
from app.services.officer_service import validate_signature_image
from app.services.serializers import iso_utc


MEDIA_TYPES = {'DOCX': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'PDF': 'application/pdf'}


def _image_bytes(value: str) -> bytes:
    validate_signature_image(value)
    return base64.b64decode(value.split(',', 1)[1])


def _record_time(value) -> str:
    if not isinstance(value, str) or not value:
        return '未记录'
    try:
        stamp = datetime.fromisoformat(value.replace('Z', '+00:00'))
        if stamp.tzinfo is None:
            stamp = stamp.replace(tzinfo=timezone.utc)
        return stamp.astimezone(timezone(timedelta(hours=8))).strftime('%Y-%m-%d %H:%M:%S')
    except ValueError:
        return '未记录'


def _paragraphs(content: dict) -> list[str]:
    """Use the frozen editor's canonical answers, with older snapshot compatibility."""
    case = content.get('case') or {}
    session = content.get('session') or {}
    facts = {item.get('key'): str(item.get('value') or '').strip() for item in content.get('facts') or []}

    def fact(key: str) -> str:
        value = facts.get(key)
        return value if value and value not in {'未录入', '未填写'} else '未记录'

    interrogator = case.get('interrogator') or {}
    recorder = case.get('recorder') or {}
    interrogator_name = interrogator.get('displayName') or interrogator.get('name') or case.get('officerName') or '未记录'
    recorder_name = recorder.get('displayName') or recorder.get('name') or fact('recorder_name')
    lines = ['询问笔录',
             f"询问时间（北京时间）：{_record_time(session.get('startedAt'))} 至 {_record_time(session.get('endedAt'))}",
             f"询问地点：{fact('interrogation_place')}",
             f'询问人：{interrogator_name}　记录人：{recorder_name}',
             f"被询问人：{case.get('suspectName') or '未记录'}　证件号码：{case.get('idNumber') or case.get('id_number') or '未记录'}",
             f"案件编号：{case.get('case_id') or case.get('id') or case.get('caseId') or '未记录'}"]
    transcript = content.get('transcript') or {}
    if transcript.get('source') == 'LEGACY_MESSAGES':
        for message in transcript.get('messages') or []:
            speaker = str(message.get('speaker') or message.get('from') or '').upper()
            prefix = '问' if speaker in {'民警', 'OFFICER', 'POLICE', 'QUESTION', 'INTERROGATOR', 'RECORDER'} else '答'
            lines.append(f"{prefix}：{message.get('text') or ''}")
    elif 'questions' in transcript:
        for question in transcript.get('questions') or []:
            if question.get('active') is False:
                continue
            answer = question.get('formalAnswerText')
            lines.extend([f"问：{question.get('text') or ''}", f"答：{answer or ''}"])
    else:
        for entry in transcript.get('entries') or []:
            lines.extend([f"问：{entry.get('formalQuestionText') or entry.get('questionText') or ''}",
                          f"答：{entry.get('answerText') or ''}"])
    return lines


def _write_docx(path: Path, lines: list[str], snapshot) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.5)
    section.left_margin = section.right_margin = Cm(2.5)
    normal = document.styles['Normal']
    normal.font.name = 'SimSun'
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.line_spacing = 1.5
    for index, text in enumerate(lines):
        paragraph = document.add_paragraph(text)
        if text.startswith('问：') or index == 0:
            paragraph.paragraph_format.keep_with_next = True
        if index == 0:
            paragraph.alignment = 1
            paragraph.runs[0].bold = True
            paragraph.runs[0].font.size = Pt(18)
    document.add_paragraph(f'冻结版本：{snapshot.version}　正本校验值：{snapshot.content_hash}')
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    document.core_properties.title = '询问笔录（纯文字版）'
    document.core_properties.author = ''
    document.save(path)


def _write_pdf(path: Path, lines: list[str], snapshot, signatures: list, fingerprint) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer

    if 'STSong-Light' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    body = ParagraphStyle('body', fontName='STSong-Light', fontSize=12, leading=21,
                          wordWrap='CJK', spaceAfter=6)
    title = ParagraphStyle('title', parent=body, fontSize=18, leading=28, alignment=TA_CENTER, spaceAfter=18)
    question = ParagraphStyle('question', parent=body, keepWithNext=True)
    small = ParagraphStyle('small', parent=body, fontSize=8, leading=12, textColor=colors.HexColor('#444444'))
    text = lambda value: escape(str(value)).replace('\n', '<br/>')
    story = [Paragraph(text(line), title if i == 0 else question if line.startswith('问：') else body)
             for i, line in enumerate(lines)]
    story.extend([Spacer(1, 12), Paragraph(f'冻结版本：{snapshot.version}', small),
                  Paragraph(f'正本 SHA-256：{snapshot.content_hash}', small)])
    marks = []
    for signature in signatures:
        role = '被询问人' if signature.signer_role == 'SUSPECT' else '民警'
        image = Image(io.BytesIO(_image_bytes(signature.image_data)))
        image._restrictSize(140, 55)
        image.hAlign = 'LEFT'
        marks.extend([Paragraph(text(f'{role}签名：{signature.signer_name}'), body), image,
                      Paragraph(text(f'签署时间：{iso_utc(signature.created_at)}'), small), Spacer(1, 8)])
    fingerprint_image = Image(io.BytesIO(_image_bytes(fingerprint.image_data)))
    fingerprint_image._restrictSize(65, 65)
    fingerprint_image.hAlign = 'LEFT'
    marks.extend([Paragraph('指纹捺印：', body), fingerprint_image])
    story.append(KeepTogether(marks))

    def footer(canvas, doc):
        canvas.setFont('STSong-Light', 9)
        canvas.drawCentredString(A4[0] / 2, 28, f'第 {doc.page} 页')

    SimpleDocTemplate(str(path), pagesize=A4, rightMargin=65, leftMargin=65,
                      topMargin=60, bottomMargin=50, title='询问笔录（签章版）', author='').build(
                          story, onFirstPage=footer, onLaterPages=footer)


class DocumentExportService:
    def __init__(self, db: Session, data_dir: Path):
        self.db = db
        self.data_dir = Path(data_dir)

    @staticmethod
    def _actor():
        actor = current_actor()
        if actor is None:
            raise DomainError('AUTH_SESSION_REQUIRED', '导出或调阅笔录须先验证民警身份', 401)
        return actor

    def _root(self) -> Path:
        root = self.data_dir.resolve()
        export_root = root / 'exports'
        export_root.mkdir(parents=True, exist_ok=True)
        if export_root.resolve() != export_root:
            raise DomainError('EXPORT_PATH_INVALID', '导出目录不能是符号链接', 409)
        return export_root

    @staticmethod
    def _payload(row: DocumentExport) -> dict:
        return {'id': row.id, 'caseId': row.case_id, 'snapshotId': row.snapshot_id, 'snapshotHash': row.snapshot_hash,
                'format': row.format, 'fileName': f'笔录-{row.id}.{row.format.lower()}',
                'sha256': row.sha256, 'sizeBytes': row.size_bytes, 'actorId': row.actor_id,
                'createdAt': iso_utc(row.created_at), 'signatureCount': len(json.loads(row.signature_manifest_json)),
                'fingerprintSha256': row.fingerprint_sha256,
                'downloadUrl': f'/api/v1/cases/{row.case_id}/exports/{row.id}/file'}

    def _count(self, case_id: str) -> int:
        return int(self.db.scalar(select(func.count()).select_from(DocumentExport).where(DocumentExport.case_id == case_id)) or 0)

    def list(self, case_id: str, *, limit: int = 100, offset: int = 0) -> dict:
        self._actor()
        assert_case_not_void(self.db, case_id)
        rows = self.db.scalars(select(DocumentExport).where(DocumentExport.case_id == case_id)
                               .order_by(DocumentExport.created_at.desc(), DocumentExport.id.desc())
                               .limit(max(1, min(200, limit))).offset(max(0, offset)))
        return {'caseId': case_id, 'exportCount': self._count(case_id), 'items': [self._payload(row) for row in rows]}

    def create(self, case_id: str, format: str) -> dict:
        actor = self._actor()
        if format not in MEDIA_TYPES:
            raise DomainError('EXPORT_FORMAT_INVALID', '导出格式仅支持 DOCX 或 PDF', 400)
        begin_sqlite_immediate(self.db)
        case = assert_case_not_void(self.db, case_id)
        snapshot = document_repo.latest_snapshot(self.db, case_id)
        if snapshot is None or case.workflow_state not in {'FROZEN', 'SIGNED', 'REPORT_GENERATED'}:
            raise DomainError('DOCUMENT_NOT_FROZEN', '请先冻结正本后再导出', 409)
        if hashlib.sha256(snapshot.content_json.encode('utf-8')).hexdigest() != snapshot.content_hash:
            raise DomainError('DOCUMENT_INTEGRITY_FAILED', '冻结正本校验失败，不能导出', 409)
        try:
            content = json.loads(snapshot.content_json)
            lines = _paragraphs(content)
        except (ValueError, TypeError, AttributeError):
            raise DomainError('DOCUMENT_CONTENT_INVALID', '冻结正本内容无效', 409)
        signatures, fingerprint, manifest = [], None, []
        if format == 'PDF':
            from app.database.document_mark_models import DocumentFingerprint

            signatures = list(self.db.scalars(select(SignatureRecord).where(
                SignatureRecord.case_id == case_id, SignatureRecord.snapshot_id == snapshot.id,
                SignatureRecord.status == 'SAVED').order_by(SignatureRecord.signer_role.desc())))
            if case.workflow_state not in {'SIGNED', 'REPORT_GENERATED'} or {row.signer_role for row in signatures} != {'SUSPECT', 'OFFICER'}:
                raise DomainError('DOCUMENT_SIGNATURES_REQUIRED', '双方签署完成后才能导出签章版 PDF', 409)
            fingerprint = self.db.scalar(select(DocumentFingerprint).where(
                DocumentFingerprint.case_id == case_id, DocumentFingerprint.snapshot_id == snapshot.id))
            if fingerprint is None:
                raise DomainError('DOCUMENT_FINGERPRINT_REQUIRED', '请先完成本版正本的指纹捺印', 409)
            if hashlib.sha256(_image_bytes(fingerprint.image_data)).hexdigest() != fingerprint.sha256:
                raise DomainError('FINGERPRINT_INTEGRITY_FAILED', '指纹图像校验失败', 409)
            manifest = [{'role': row.signer_role, 'name': row.signer_name,
                         'sha256': hashlib.sha256(_image_bytes(row.image_data)).hexdigest()} for row in signatures]
        export_id = str(uuid4())
        stored_name = f'{export_id}.{format.lower()}'
        path = self._root() / stored_name
        temporary = path.with_suffix('.tmp')
        try:
            if format == 'DOCX':
                _write_docx(temporary, lines, snapshot)
            else:
                _write_pdf(temporary, lines, snapshot, signatures, fingerprint)
            temporary.replace(path)
            row = DocumentExport(id=export_id, case_id=case_id, snapshot_id=snapshot.id,
                                 snapshot_hash=snapshot.content_hash, format=format, stored_name=stored_name,
                                 sha256=hashlib.sha256(path.read_bytes()).hexdigest(), size_bytes=path.stat().st_size,
                                 actor_id=actor.badge_number, signature_manifest_json=json.dumps(manifest, ensure_ascii=False),
                                 fingerprint_sha256=fingerprint.sha256 if fingerprint else None)
            self.db.add(row)
            self.db.flush()
            count = self._count(case_id)
            audit_repo.add(self.db, case_id=case_id, actor_id=actor.badge_number, action='DOCUMENT_EXPORT',
                           target_type='DOCUMENT_EXPORT', target_id=row.id,
                           after={'format': format, 'snapshotId': snapshot.id, 'snapshotHash': snapshot.content_hash,
                                  'sha256': row.sha256, 'exportCount': count})
            self.db.commit()
            return {**self._payload(row), 'exportCount': count}
        except Exception:
            self.db.rollback()
            temporary.unlink(missing_ok=True)
            path.unlink(missing_ok=True)
            raise

    def download(self, case_id: str, export_id: str) -> tuple[Path, dict]:
        actor = self._actor()
        assert_case_not_void(self.db, case_id)
        row = self.db.get(DocumentExport, export_id)
        if row is None or row.case_id != case_id:
            raise DomainError('EXPORT_NOT_FOUND', '未找到本案导出记录', 404)
        if row.format not in MEDIA_TYPES or row.stored_name != f'{row.id}.{row.format.lower()}' or Path(row.stored_name).name != row.stored_name:
            raise DomainError('EXPORT_PATH_INVALID', '导出文件路径无效', 409)
        root = self._root()
        path = root / row.stored_name
        if path.resolve().parent != root or not path.is_file():
            raise DomainError('EXPORT_FILE_UNAVAILABLE', '导出文件不存在或路径异常', 409)
        if path.stat().st_size != row.size_bytes or hashlib.sha256(path.read_bytes()).hexdigest() != row.sha256:
            raise DomainError('EXPORT_INTEGRITY_FAILED', '导出文件校验失败', 409)
        audit_repo.add(self.db, case_id=case_id, actor_id=actor.badge_number, action='DOCUMENT_EXPORT_DOWNLOAD',
                       target_type='DOCUMENT_EXPORT', target_id=row.id,
                       detail={'format': row.format, 'sha256': row.sha256})
        self.db.commit()
        return path, self._payload(row)
